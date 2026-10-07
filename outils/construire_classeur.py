"""Construit Prestations_CDamien_8.xlsm à partir :
  - du prototype Prestations_CDamien_7.xlsm (onglets jours, réglages, distances, modèle de carte, projet VBA d'origine)
  - de l'export d'abonnements Abonnements_iGclean071026.xlsx (onglet Source)
  - des macros du dossier vba/

usage : python3 outils/construire_classeur.py [sortie.xlsm]
Le classeur produit a ses onglets calculés vides : ils se remplissent à l'ouverture (macros activées),
ou avec outils/remplir_classeur.py (exécution des macros dans LibreOffice).
"""
import datetime as dt
import glob
import os
import re
import sys
import tempfile
import zipfile

import olefile
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from oletools.olevba import decompress_stream

ICI = os.path.dirname(os.path.abspath(__file__))
RACINE = os.path.dirname(ICI)
sys.path.insert(0, os.path.join(ICI, 'vbaproject'))
from buildvba import build  # noqa: E402
from package import make_xlsm  # noqa: E402

PROTO = os.path.join(RACINE, 'Prestations_CDamien_7.xlsm')
ABOS = os.path.join(RACINE, 'Abonnements_iGclean071026.xlsx')
args = [x for x in sys.argv[1:] if not x.startswith('--remplir=')]
REMPLIR = next((x.split('=', 1)[1] for x in sys.argv[1:] if x.startswith('--remplir=')), None)
SORTIE = args[0] if args else os.path.join(RACINE, 'Prestations_CDamien_8.xlsm')

BLEU = '1F3A5F'
GRIS_CLAIR = 'E7ECF2'
FOND_TUILE = 'F4F6F9'
SAISIE = 'FFF8E1'
POLICE = 'Arial'

SYSTEME = ['DASHBOARD', 'CONTROLES', 'PARAMETRES', 'ABONNEMENTS', 'DATA', 'ABOS', 'LISTES', 'TOURNEES',
           'DISTANCES', 'MODELE_CARTE']


def f(size=10, bold=False, color='1E1E1E', italic=False, underline=None):
    return Font(name=POLICE, size=size, bold=bold, color=color, italic=italic, underline=underline)


def fill(c):
    return PatternFill('solid', start_color=c, end_color=c)


def bouton(ws, ref, texte, plage=None, clair=False):
    c = ws[ref]
    c.value = texte
    c.hyperlink = Hyperlink(ref=ref, location="'%s'!%s" % (ws.title, ref), display=texte)
    c.font = f(10, True, BLEU if clair else 'FFFFFF')
    c.fill = fill(GRIS_CLAIR if clair else BLEU)
    c.alignment = Alignment(horizontal='center', vertical='center')
    if plage:
        ws.merge_cells(plage)


def bande(ws, plage, texte):
    first = plage.split(':')[0]
    ws[first].value = texte
    for row in ws[plage]:
        for c in row:
            c.fill = fill(BLEU)
            c.font = f(11, True, 'FFFFFF')
            c.alignment = Alignment(vertical='center')
    ws.row_dimensions[ws[first].row].height = 20


def entete(ws, plage, textes):
    for c, t in zip(ws[plage][0], textes):
        c.value = t
        c.font = f(10, True, BLEU)
        c.fill = fill(GRIS_CLAIR)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = Border(bottom=Side(style='thin', color=BLEU))


def largeurs(ws, d):
    for col, w in d.items():
        ws.column_dimensions[col].width = w


def oui_non(ws, plage):
    dv = DataValidation(type='list', formula1='"Oui,Non"', allow_blank=True, showErrorMessage=True)
    dv.error = 'Oui ou Non'
    dv.add(plage)
    ws.add_data_validation(dv)


# ---------------------------------------------------------------------------
proto_v = openpyxl.load_workbook(PROTO, data_only=True)
abos_v = openpyxl.load_workbook(ABOS, data_only=True)

wb = openpyxl.Workbook()
wb._named_styles['Normal'].font = Font(name=POLICE, size=10)
ws_dash = wb.active
ws_dash.title = 'DASHBOARD'
ws_ctrl = wb.create_sheet('CONTROLES')
ws_par = wb.create_sheet('PARAMETRES')
ws_abo = wb.create_sheet('ABONNEMENTS')

# --- onglets jours (tels qu'exportés), nommés d'après la date de leurs lignes.
# Même règle que la macro RenommerOngletsJours : un jour présent deux fois -> la version
# la plus à droite est gardée, l'autre devient « JJMMAA ancien ».
def dates_onglet(src):
    ds = set()
    for r in src.iter_rows(values_only=True):
        if r and isinstance(r[0], dt.datetime) and len(r) > 6 and str(r[6] or '').strip():
            ds.add(r[0].date())
    return ds


jours_proto = [n for n in proto_v.sheetnames if re.fullmatch(r'\d{6}', n)]
nom_final, porteur, plus_recent = {}, {}, {}
for nom in jours_proto:
    ds = dates_onglet(proto_v[nom])
    cible = min(ds).strftime('%d%m%y') + ('' if len(ds) == 1 else '-' + max(ds).strftime('%d%m%y'))
    if cible in porteur:
        nom_final[porteur[cible]] = cible + ' ancien'
    porteur[cible] = nom
    nom_final[nom] = cible
    plus_recent[nom] = max(ds)
gardes = sorted([n for n in jours_proto if not nom_final[n].endswith('ancien')], key=lambda n: plus_recent[n], reverse=True)
ecartes = [n for n in jours_proto if nom_final[n].endswith('ancien')]
jours = []
for nom in gardes + ecartes:
    src = proto_v[nom]
    ws = wb.create_sheet(nom_final[nom])
    jours.append(nom_final[nom])
    for row in src.iter_rows():
        for c in row:
            if c.value is None:
                continue
            d = ws.cell(row=c.row, column=c.column, value=c.value)
            if isinstance(c.value, dt.datetime):
                d.number_format = 'dd/mm/yyyy'
            elif isinstance(c.value, dt.time):
                d.number_format = 'hh:mm'
    for col in range(1, 32):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 12
    ws.column_dimensions['G'].width = 24
    ws.column_dimensions['P'].width = 28
    ws.column_dimensions['Q'].width = 34
    ws.freeze_panes = 'A2'
    ws.sheet_properties.tabColor = 'A6A6A6' if nom in ecartes else '2E7D32'

# --- techniques (masqués)
ws_data = wb.create_sheet('DATA')
ws_aboss = wb.create_sheet('ABOS')
ws_l = wb.create_sheet('LISTES')
ws_t = wb.create_sheet('TOURNEES')
ws_di = wb.create_sheet('DISTANCES')
ws_m = wb.create_sheet('MODELE_CARTE')
for ws in (ws_data, ws_aboss, ws_l, ws_t, ws_di, ws_m):
    ws.sheet_state = 'hidden'

# ===========================================================================
# DASHBOARD
# ===========================================================================
ws = ws_dash
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = BLEU
largeurs(ws, {'A': 2, 'B': 34, **{c: 11.5 for c in 'CDEFGHIJKLMN'}, 'O': 12, **{c: 10.5 for c in 'PQRSTU'}})
ws['B1'].value = "Prestations C'Damien"
ws['B1'].font = f(18, True, BLEU)
ws.row_dimensions[1].height = 30
bouton(ws, 'I1', 'Mettre à jour', 'I1:J1')
bouton(ws, 'K1', 'Ouvrir la carte', 'K1:M1')
ws['B2'].value = ("Activez les macros (bandeau « Activer le contenu ») : le tableau de bord se construit "
                  "automatiquement à partir des onglets sources.")
ws['B2'].font = f(9, color='6E6E6E', italic=True)
ws['B3'].font = f(10, True, 'C62828')
ws.row_dimensions[4].height = 8
ws['B5'].value = 'Agent'
ws['G5'].value = 'Du'
ws['I5'].value = 'Au'
for ref in ('B5', 'G5', 'I5'):
    ws[ref].font = f(10, True, BLEU)
    ws[ref].alignment = Alignment(horizontal='right', vertical='center')
ws.merge_cells('C5:E5')
ws['C5'].value = 'Tous'
cadre = Border(*(Side(style='thin', color='9FB3C8'),) * 4)
for ref in ('C5', 'D5', 'E5', 'H5', 'J5'):
    ws[ref].border = cadre
    ws[ref].fill = fill('FFFFFF')
for ref in ('C5', 'H5', 'J5'):
    ws[ref].font = f(10, True)
    ws[ref].alignment = Alignment(horizontal='center', vertical='center')
ws['H5'].number_format = 'dd/mm/yyyy'
ws['J5'].number_format = 'dd/mm/yyyy'
for ref, t in (('K5', 'Dernier jour'), ('L5', '7 jours'), ('M5', '30 jours'), ('N5', 'Tout')):
    bouton(ws, ref, t, clair=True)
ws.row_dimensions[5].height = 22
ws.row_dimensions[6].height = 8
dv = DataValidation(type='list', formula1='ListeAgents', allow_blank=False, showErrorMessage=True)
dv.error = "Choisir un agent dans la liste (ou « Tous »)."
dv.errorTitle = 'Agent'
dv.add('C5')
ws.add_data_validation(dv)
dvd = DataValidation(type='date', operator='greaterThan', formula1='36526', showErrorMessage=True)
dvd.error = 'Saisir une date (jj/mm/aaaa).'
dvd.add('H5')
dvd.add('J5')
ws.add_data_validation(dvd)

tuiles = [('B', 'B', 'Prestations', '0'), ('C', 'D', 'Pointées', '0%'), ('E', 'F', 'Heures prévues', '0.0'),
          ('G', 'H', 'Heures pointées', '0.0'), ('I', 'J', 'Chantiers servis', '0'),
          ('K', 'L', 'CA mensuel', '#,##0 "€"'), ('M', 'N', 'Km pointés', '#,##0')]
for c1, c2, lib, fmt in tuiles:
    for r in (7, 8, 9):
        if c1 != c2:
            ws.merge_cells('%s%d:%s%d' % (c1, r, c2, r))
        for col in sorted({c1, c2}):
            ws['%s%d' % (col, r)].fill = fill(FOND_TUILE)
    ws['%s7' % c1].value = lib.upper()
    ws['%s7' % c1].font = f(8, True, '6E7F91')
    ws['%s8' % c1].font = f(18, True, BLEU)
    ws['%s8' % c1].number_format = fmt
    ws['%s9' % c1].font = f(8, color='6E7F91')
    for r in (7, 8, 9):
        ws['%s%d' % (c1, r)].alignment = Alignment(horizontal='center', vertical='center')
    ws['%s7' % c1].border = Border(top=Side(style='medium', color=BLEU))
    if c1 != c2:
        ws['%s7' % c2].border = Border(top=Side(style='medium', color=BLEU))
ws.row_dimensions[7].height = 18
ws.row_dimensions[8].height = 30
ws.row_dimensions[9].height = 16
ws.freeze_panes = 'A10'

# ===========================================================================
# CONTROLES
# ===========================================================================
ws = ws_ctrl
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = 'E67800'
largeurs(ws, {'A': 3, 'B': 16, 'C': 40, 'D': 13, 'E': 13, 'F': 34, 'G': 13, 'H': 13, 'I': 46, 'J': 16})
ws.column_dimensions['A'].hidden = True
ws['B1'].value = 'Contrôles : à corriger'
ws['B1'].font = f(18, True, BLEU)
ws.row_dimensions[1].height = 30
bouton(ws, 'I1', 'Mettre à jour', 'I1:J1')

# ===========================================================================
# PARAMETRES
# ===========================================================================
ws = ws_par
ws.sheet_view.showGridLines = False
ws.sheet_properties.tabColor = '7F7F7F'
largeurs(ws, {'A': 2, 'B': 34, 'C': 16, 'D': 44, 'E': 14, 'F': 14, 'G': 14, 'H': 11, 'I': 11, 'J': 30, 'K': 3,
              'L': 38, 'M': 16, 'N': 34})
ws.column_dimensions['J'].hidden = True
ws['B1'].value = 'Paramètres'
ws['B1'].font = f(18, True, BLEU)
ws.row_dimensions[1].height = 30
bande(ws, 'B3:N3', "MODE D'EMPLOI")
aide = [
    "1. Abonnements : copier l'export des abonnements et le coller dans l'onglet ABONNEMENTS, cellule A1 (remplacer tout).",
    "2. Validations du jour : ajouter un onglet (bouton +) et y coller l'extract en A1. Son nom n'a pas d'importance : "
    "il est renommé automatiquement à la date des prestations (ex. 071026).",
    "3. Revenir sur DASHBOARD : tout se met à jour tout seul (sinon, bouton « Mettre à jour »).",
    "4. CONTROLES liste ce qui est à corriger : chantiers sans abonnement, abonnements sans prestation, prestations non "
    "affectées, pointages incomplets, jours manquants...",
    "5. Ci-dessous : adresse de l'agence, agents comptés ou non dans les indicateurs, domiciles et passages à l'agence "
    "(pour les km), rattachements manuels. Les cellules jaunes se modifient.",
]
for i, t in enumerate(aide):
    ws['B%d' % (4 + i)].value = t
    ws['B%d' % (4 + i)].font = f(10)
bande(ws, 'B10:I10', 'AGENCE ET RÉGLAGES')
ws['B11'].value = "Adresse de l'agence"
ws.merge_cells('C11:G11')
ws['B12'].value = "Retour à l'agence le soir"
ws['B13'].value = 'Distance max. agence - chantier (km)'
ws['B14'].value = 'Période de contrôle (jours)'
ws['D12'].value = "Pour les agents sans domicile renseigné : fin de tournée à l'agence (km)."
ws['D13'].value = "Au-delà, les coordonnées GPS d'un chantier sont considérées comme fausses."
ws['D14'].value = 'Derniers jours de données pris pour les contrôles et la part de CA de chaque agent.'
for r in (12, 13, 14):
    ws['D%d' % r].font = f(9, color='6E6E6E', italic=True)
ws['H10'].value = 'Latitude'
ws['I10'].value = 'Longitude'
for ref in ('B11', 'B12', 'B13', 'B14'):
    ws[ref].font = f(10, True)
for ref in ('C11', 'C12', 'C13', 'C14'):
    ws[ref].fill = fill(SAISIE)
    ws[ref].border = cadre
    ws[ref].alignment = Alignment(horizontal='left' if ref == 'C11' else 'center')
for col in 'DEFG':
    ws['%s11' % col].fill = fill(SAISIE)
    ws['%s11' % col].border = cadre
oui_non(ws, 'C12')
bouton(ws, 'B16', 'Localiser les adresses et recalculer les km', 'B16:C16')
ws.row_dimensions[16].height = 22
bande(ws, 'B18:I18', 'AGENTS')
bande(ws, 'L18:N18', 'RATTACHEMENTS MANUELS')
entete(ws, 'B19:J19', ['Agent', 'Compté dans les indicateurs', 'Adresse du domicile', 'Départ du domicile',
                       'Passage agence le matin', 'Passage agence le soir', 'Latitude', 'Longitude',
                       'Adresse localisée'])
entete(ws, 'L19:N19', ['Projet ou chantier', 'N° abonnement', 'Statut'])
ws.row_dimensions[19].height = 30
for r in range(20, 420):
    for col in 'CDEFG':
        ws['%s%d' % (col, r)].fill = fill(SAISIE)
    for col in 'CEFG':
        ws['%s%d' % (col, r)].alignment = Alignment(horizontal='center')
    for col in 'LM':
        ws['%s%d' % (col, r)].fill = fill(SAISIE)
oui_non(ws, 'C20:C420')
oui_non(ws, 'E20:G420')
ws.freeze_panes = 'A20'

# reprise des réglages du prototype
pp = proto_v['PARAMETRES']
pl = proto_v['LISTES']
ws['C11'].value = pp['C4'].value
ws['H11'].value = pp['F4'].value
ws['I11'].value = pp['G4'].value
ws['J11'].value = pp['H4'].value
ws['C12'].value = 'Oui' if str(pp['C6'].value or '').strip().lower() == 'oui' else 'Non'
ws['C13'].value = pp['C7'].value or 60
ws['C14'].value = pp['K6'].value or 28
inclus, matin, soir = {}, {}, {}
for row in pl.iter_rows(min_row=3, values_only=True):
    if row[0]:
        nom = str(row[0]).strip()
        inclus[nom] = row[2]
        matin[nom] = row[5]
        soir[nom] = row[6]


def fictif(nom):
    n = nom.upper()
    return any(k in n for k in ('SUPPRIMER', 'REMPLACEMENT', 'DEFINIR', 'NOUVEAU'))


agents = []
for row in pp.iter_rows(min_row=13, values_only=True):
    nom = str(row[1] or '').strip()
    if not nom:
        continue
    compte = inclus.get(nom)
    compte = ('Non' if fictif(nom) else 'Oui') if compte is None else ('Oui' if compte == 1 else 'Non')
    adr = row[2]
    depart = 'Oui' if adr and row[3] == '☑' else 'Non'
    agents.append((nom, compte, adr, depart, 'Oui' if matin.get(nom) == 1 else 'Non',
                   'Oui' if soir.get(nom) == 1 else 'Non', row[5] if adr else None, row[6] if adr else None,
                   row[7] if adr else None))
agents.sort(key=lambda a: a[0].lower())
for i, a in enumerate(agents):
    for j, val in enumerate(a):
        ws.cell(row=20 + i, column=2 + j, value=val)

# ===========================================================================
# ABONNEMENTS (export collé tel quel)
# ===========================================================================
ws = ws_abo
ws.sheet_properties.tabColor = '00838F'
src = abos_v['Source']
for row in src.iter_rows():
    for c in row:
        if c.value is None:
            continue
        d = ws.cell(row=c.row, column=c.column, value=c.value)
        if isinstance(c.value, dt.datetime):
            d.number_format = 'dd/mm/yyyy'
for c in ws[1]:
    c.font = f(10, True)
largeurs(ws, {'A': 7, 'B': 11, 'C': 11, 'D': 11, 'E': 15, 'F': 9, 'G': 34, 'H': 30, 'I': 30, 'J': 4, 'K': 7,
              'L': 16, 'M': 30, 'N': 11, 'O': 10, 'P': 10, 'Q': 10, 'R': 5, 'S': 12})
ws.freeze_panes = 'A2'

# ===========================================================================
# Techniques
# ===========================================================================
ws_l['A1'].value = 'Agents'
ws_l['A2'].value = 'Tous'
ws_l['I1'].value = 'Clé'
ws_l['J1'].value = 'Valeur'
ws_t['A1'].value = 'Tournées : une ligne par agent et par jour (km pointés, dans l\'ordre des clics)'
for r, row in enumerate(proto_v['DISTANCES'].iter_rows(values_only=True), start=1):
    for c, val in enumerate(row[:3], start=1):
        if val is not None:
            ws_di.cell(row=r, column=c, value=val)
# modèle de carte : les km du planning ne sont plus calculés -> on ne les affiche que s'ils existent
html = ''.join(str(proto_v['MODELE_CARTE'].cell(i, 1).value or '') for i in range(1, 501))
avant = "$('stats').innerHTML += ' \\u00b7 <b>' + kmTot(t[0]) + '</b> planning \\u00b7 <b>' + kmTot(t[1]) + '</b> point\\u00e9s' +"
apres = "$('stats').innerHTML += (t[0] ? ' \\u00b7 <b>' + kmTot(t[0]) + '</b> planning' : '') + ' \\u00b7 <b>' + kmTot(t[1]) + '</b> point\\u00e9s' +"
assert avant in html, 'modèle de carte : ligne des km introuvable'
html = html.replace(avant, apres)
TAILLE = 30000
for i in range(0, len(html), TAILLE):
    ws_m.cell(row=i // TAILLE + 1, column=1, value=html[i:i + TAILLE])

wb.defined_names['ListeAgents'] = DefinedName('ListeAgents', attr_text="'LISTES'!$A$2:$A$3")
wb.active = 0

if REMPLIR:
    from copy import copy

    lo = openpyxl.load_workbook(REMPLIR)
    assert [n for n in lo.sheetnames] == wb.sheetnames, (lo.sheetnames, wb.sheetnames)

    def copier(src, dst, style=True):
        dst.value = src.value
        if src.has_style:
            dst.number_format = src.number_format
            if style:
                dst.font = copy(src.font)
                dst.fill = copy(src.fill)
                dst.border = copy(src.border)
                dst.alignment = copy(src.alignment)

    for nom in ('DATA', 'ABOS', 'LISTES', 'TOURNEES', 'DISTANCES'):
        for row in lo[nom].iter_rows():
            for c in row:
                if c.value is not None:
                    copier(c, wb[nom].cell(row=c.row, column=c.column), style=False)
    for nom, l1 in (('CONTROLES', 2), ('DASHBOARD', 11)):
        src, dst = lo[nom], wb[nom]
        for row in src.iter_rows(min_row=l1):
            for c in row:
                if c.value is not None or c.has_style:
                    copier(c, dst.cell(row=c.row, column=c.column))
        for r, d in src.row_dimensions.items():
            if r >= l1 and d.height:
                dst.row_dimensions[r].height = d.height
    d, s_ = lo['DASHBOARD'], wb['DASHBOARD']
    for ref in ('B2', 'B3', 'C5', 'H5', 'J5', 'B8', 'C8', 'E8', 'G8', 'I8', 'K8', 'M8', 'B9', 'C9', 'E9', 'G9', 'I9',
                'K9', 'M9'):
        s_[ref].value = d[ref].value
    s_['B3'].font = copy(d['B3'].font)
    s_['B2'].font = f(9, color='6E6E6E', italic=True)
    p_lo, p_ = lo['PARAMETRES'], wb['PARAMETRES']
    for row in p_lo.iter_rows(min_row=11, max_row=11, min_col=8, max_col=10):
        for c in row:
            p_.cell(row=c.row, column=c.column).value = c.value
    for row in p_lo.iter_rows(min_row=20, min_col=2, max_col=14):
        for c in row:
            p_.cell(row=c.row, column=c.column).value = c.value
    nb_agents = sum(1 for v in lo['LISTES']['A'] if v.value not in (None, ''))
    wb.defined_names['ListeAgents'] = DefinedName('ListeAgents', attr_text="'LISTES'!$A$2:$A$%d" % (nb_agents))

# ===========================================================================
# Macros
# ===========================================================================
codenames = {'DASHBOARD': 'shDashboard', 'CONTROLES': 'shControles', 'PARAMETRES': 'shParametres',
             'ABONNEMENTS': 'shAbonnements', 'DATA': 'shData', 'ABOS': 'shAbos', 'LISTES': 'shListes',
             'TOURNEES': 'shTournees', 'DISTANCES': 'shDistances', 'MODELE_CARTE': 'shModele'}
for nom in jours:
    codenames[nom] = 'shJ' + re.sub(r'\W', '_', nom)
VBA = os.path.join(RACINE, 'vba')


def lire(nom):
    p = os.path.join(VBA, nom)
    return open(p, encoding='utf-8').read() if os.path.exists(p) else 'Option Explicit\n'


modules = [('ThisWorkbook', 'doc', lire('ThisWorkbook.cls'), 'wb')]
for onglet in wb.sheetnames:
    cn = codenames[onglet]
    modules.append((cn, 'doc', lire(cn + '.cls'), 'ws'))
for nom in ('Outils', 'Pipeline', 'Abonnements', 'Dashboard', 'Controles', 'Carte'):
    modules.append((nom, 'std', lire(nom + '.bas'), ''))

tmp = tempfile.mkdtemp()
z = zipfile.ZipFile(PROTO)
vba_bin = os.path.join(tmp, 'orig.bin')
open(vba_bin, 'wb').write(z.read('xl/vbaProject.bin'))
o = olefile.OleFileIO(vba_bin)
projet = o.openstream('PROJECT').read().decode('cp1252')
dir_orig = decompress_stream(bytearray(o.openstream('VBA/dir').read()))
o.close()

xlsx = os.path.join(tmp, 'classeur.xlsx')
wb.save(xlsx)
nouveau = os.path.join(tmp, 'vbaProject.bin')
build(modules, bytes(dir_orig), projet, nouveau)
make_xlsm(xlsx, nouveau, SORTIE, codenames)
print('Classeur écrit :', SORTIE)
